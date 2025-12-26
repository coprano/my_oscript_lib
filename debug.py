from openpyxl import load_workbook
import os

wb = load_workbook("your_input_file.xlsx")
for sheet_name in wb.sheetnames:
    ws = wb[sheet_name]
    print(f"\nSheet: {sheet_name}")
    print(f"  _images: {getattr(ws, '_images', 'NOT FOUND')}")
    print(f"  _drawings: {getattr(ws, '_drawings', 'NOT FOUND')}")
    
    # Check if images exist in different way
    if hasattr(ws, '_images') and ws._images:
        for i, img in enumerate(ws._images):
            print(f"    Image {i}: {type(img)}, anchor={getattr(img, 'anchor', None)}")
    
    if hasattr(ws, '_drawings') and ws._drawings:
        for i, drawing in enumerate(ws._drawings):
            print(f"    Drawing {i}: {type(drawing)}")
            if hasattr(drawing, '_charts'):
                print(f"      Charts: {len(drawing._charts)}")
            if hasattr(drawing, 'shapes'):
                print(f"      Shapes: {len(drawing.shapes)}")

# Also check the file structure
print("\n\nZip contents (xl/media):")
import zipfile
with zipfile.ZipFile("your_input_file.xlsx", 'r') as z:
    for name in z.namelist():
        if 'media' in name or 'drawing' in name:
            print(f"  {name}")